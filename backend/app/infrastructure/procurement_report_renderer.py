"""Render the same immutable report snapshot to PDF or XLSX."""

import re
from decimal import Decimal
from html import escape
from io import BytesIO

from app.domain.entities.tabular_report import TabularReport


def _text(value: str) -> str:
    return re.sub(r"[\x00-\x08\x0b-\x1f]", " ", value)


def _pdf_metadata(report, width, paragraph):
    from reportlab.platypus import LongTable, TableStyle

    rows = []
    for index in range(0, len(report.metadata), 2):
        pairs = list(report.metadata[index : index + 2])
        if len(pairs) == 1:
            pairs.append(("", ""))
        rows.append([paragraph(value) for pair in pairs for value in pair])
    table = LongTable(rows, colWidths=[90, width / 2 - 90, 90, width / 2 - 90], hAlign="LEFT")
    table.setStyle(
        TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("BOTTOMPADDING", (0, 0), (-1, -1), 4)])
    )
    return table


def render_pdf(report: TabularReport) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_RIGHT
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

    output = BytesIO()
    doc = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=28,
        leftMargin=28,
        topMargin=28,
        bottomMargin=32,
        title=f"{report.title} {report.code}",
        author=report.company_name,
    )
    styles = getSampleStyleSheet()
    body = ParagraphStyle("ReportBody", fontName="Helvetica", fontSize=8, leading=11)
    right = ParagraphStyle("ReportNumber", parent=body, alignment=TA_RIGHT)
    heading = ParagraphStyle(
        "ReportSection",
        parent=styles["Heading2"],
        fontSize=11,
        spaceBefore=14,
        spaceAfter=7,
        keepWithNext=True,
    )

    def paragraph(value, style=body):
        return Paragraph(escape(_text(str(value))).replace("\n", "<br/>"), style)

    content = [
        paragraph(report.company_name, styles["Heading2"]),
        paragraph(f"{report.title} - {report.code}", styles["Title"]),
    ]
    metadata = _pdf_metadata(report, doc.width, paragraph)
    content.extend([metadata, Spacer(1, 8)])
    for section in report.sections:
        content.append(paragraph(section.title, heading))
        data = [[paragraph(col) for col in section.columns]]
        for row in section.rows:
            cells = []
            for index, value in enumerate(row):
                if isinstance(value, Decimal):
                    number_format = section.numeric_formats.get(index, "")
                    digits = (
                        0
                        if number_format == "#,##0"
                        else 6
                        if number_format.endswith("000000")
                        else 2
                    )
                    cells.append(paragraph(f"{value:,.{digits}f}", right))
                else:
                    cells.append(paragraph(value))
            data.append(cells)
        if not section.rows:
            data.append([paragraph("Sin registros"), *[paragraph("") for _ in section.columns[1:]]])
        scale = doc.width / sum(section.widths)
        table = LongTable(
            data, colWidths=[width * scale for width in section.widths], repeatRows=1, hAlign="LEFT"
        )
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8edf2")),
                    (
                        "ROWBACKGROUNDS",
                        (0, 1),
                        (-1, -1),
                        [colors.white, colors.HexColor("#f7f9fb")],
                    ),
                    ("LINEBELOW", (0, 0), (-1, 0), 0.7, colors.HexColor("#788899")),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("TOPPADDING", (0, 0), (-1, -1), 6),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ]
            )
        )
        content.append(table)
    content.append(Spacer(1, 14))
    for note in report.notes:
        content.extend([paragraph(note), Spacer(1, 5)])

    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillColor(colors.HexColor("#526170"))
        canvas.drawString(28, 17, _text(report.code))
        canvas.drawRightString(document.pagesize[0] - 28, 17, f"Página {document.page}")
        canvas.restoreState()

    doc.build(content, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def _append_excel_row(sheet, values):
    sheet.append([_text(value) if isinstance(value, str) else value for value in values])
    # User text must stay text, even if it starts with =, +, - or @.
    for cell, value in zip(sheet[sheet.max_row], values, strict=True):
        if isinstance(value, str):
            cell.data_type = "s"


def _style_excel_sheet(sheet):
    from openpyxl.styles import Alignment, Font, PatternFill

    sheet.sheet_view.showGridLines = False
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="20364A")
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        sheet.row_dimensions[row[0].row].height = max(
            20,
            max(
                len(str(cell.value or ""))
                // max(12, int(sheet.column_dimensions[cell.column_letter].width or 15))
                + 1
                for cell in row
            )
            * 15,
        )


def _add_excel_section(book, section):
    from openpyxl.utils import get_column_letter

    sheet = book.create_sheet(section.title[:31])
    _append_excel_row(sheet, section.columns)
    for row in section.rows:
        _append_excel_row(sheet, row)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.print_title_rows = "1:1"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    for index, width in enumerate(section.widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = max(15, min(50, width * 13))
    for index, number_format in section.numeric_formats.items():
        for column in sheet.iter_cols(min_col=index + 1, max_col=index + 1, min_row=2):
            for cell in column:
                cell.number_format = number_format


def render_xlsx(report: TabularReport) -> bytes:
    from openpyxl import Workbook

    book = Workbook()
    info = book.active
    info.title = "Información"

    _append_excel_row(info, (report.title, report.code))
    _append_excel_row(info, ("Empresa", report.company_name))
    for pair in report.metadata:
        _append_excel_row(info, pair)
    _append_excel_row(info, ("Notas", ""))
    for note in report.notes:
        _append_excel_row(info, (note, ""))
    info.column_dimensions["A"].width = 56
    info.column_dimensions["B"].width = 65
    for section in report.sections:
        _add_excel_section(book, section)
    for sheet in book:
        _style_excel_sheet(sheet)
    output = BytesIO()
    book.save(output)
    return output.getvalue()
