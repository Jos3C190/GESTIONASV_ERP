"""Authorized statements, full-result downloads, and immutable weighted history."""

import uuid
from datetime import datetime, timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from app.application.inventory import InventoryUseCases
from app.domain.entities.inventory_kardex import LOCAL_TZ
from app.infrastructure.db.session import async_session_factory
from app.infrastructure.models.inventory import InventoryItemModel
from app.infrastructure.models.organization import Location, Warehouse
from app.infrastructure.repositories.inventory_repository import SqlAlchemyInventoryRepository
from openpyxl import load_workbook
from sqlalchemy import select

from tests.e2e.test_inventory_preparation import _operator_headers
from tests.e2e.test_purchase_orders import _headers, _reference_ids

pytestmark = pytest.mark.e2e


async def test_kardex_statements_history_filters_exports_and_permissions(e2e_client, tmp_path):
    client = e2e_client
    company, _, warehouse, product, _, _ = await _reference_ids()
    currency = "USD"
    headers = await _headers(client, company_id=company, username="kardex-admin", is_superuser=True)
    actor = uuid.UUID((await client.get("/api/v1/auth/me", headers=headers)).json()["id"])
    overview = (await client.get("/api/v1/inventory/valuation", headers=headers)).json()
    cut = await client.post(
        "/api/v1/inventory/valuation/opening",
        headers=headers,
        json={
            "currency": currency,
            "reference": "Corte para ejercicio Kardex",
            "approved": True,
            "snapshot_fingerprint": overview["snapshot_fingerprint"],
            "lines": [],
        },
    )
    assert cut.status_code == 200, cut.text
    async with async_session_factory() as session:
        item = await session.scalar(
            select(InventoryItemModel).where(
                InventoryItemModel.company_id == company, InventoryItemModel.product_id == product
            )
        )
        location = await session.scalar(
            select(Location).where(
                Location.warehouse_id == warehouse,
                Location.storage_eligible.is_(True),
                Location.is_active.is_(True),
                Location.lifecycle_status == "active",
            )
        )
        assert item
        assert location
        item_id, location_id = item.id, location.id
        for index, (qty, value) in enumerate((("5", "50"), ("10", "150"), ("15", "300")), 1):
            await InventoryUseCases(SqlAlchemyInventoryRepository(session)).post_movement(
                company_id=company,
                actor_id=actor,
                movement_type="receipt",
                idempotency_key=f"statement-receipt-{index}",
                source_reference=f"DEMO-INGRESO-{index}",
                lines=[
                    {
                        "inventory_item_id": item_id,
                        "to_location_id": location_id,
                        "to_stock_status": "available",
                        "quantity_base": Decimal(qty),
                        "handling_unit_code": f"HU-STATEMENT-{index}",
                        "actual_measures": {
                            "gross_weight_kg": Decimal("1"),
                            "volume_m3": Decimal("0.001"),
                            "length_m": None,
                            "width_m": None,
                            "height_m": None,
                        },
                        "valuation_currency": currency,
                        "valuation_value": Decimal(value),
                    }
                ],
            )
        await session.commit()
    today = datetime.now(LOCAL_TZ).date()
    root = "/api/v1/inventory/kardex"
    params = {
        "inventory_item_id": str(item_id),
        "period_from": today.isoformat(),
        "period_to": today.isoformat(),
    }
    physical_params = {
        **params,
        "warehouse_id": str(warehouse),
        "source_reference": "DEMO-INGRESO-2",
    }
    physical = await client.get(f"{root}/physical", headers=headers, params=physical_params)
    assert physical.status_code == 200, physical.text
    physical_data = physical.json()
    assert physical_data["total"] == 1
    assert Decimal(physical_data["incoming_quantity"]) == 10  # Same document filter as the rows.
    assert (
        Decimal(physical_data["closing_quantity"]) == 30
    )  # Whole product sequence, not filter sum.
    assert physical_data["items"][0]["operator"] == "kardex-admin"
    assert physical_data["context"]["company_name"]
    valued = await client.get(f"{root}/valuation", headers=headers, params=params)
    assert valued.status_code == 200, valued.text
    values = valued.json()
    assert [Decimal(e["average_unit_cost"]) for e in values["entries"]] == [
        Decimal("10"),
        Decimal("13.333333"),
        Decimal("16.666667"),
    ]
    assert [Decimal(e["balance_quantity"]) for e in values["entries"]] == [5, 15, 30]
    assert [Decimal(e["balance_value"]) for e in values["entries"]] == [50, 200, 500]
    assert [Decimal(e["applied_unit_cost"]) for e in values["entries"]] == [10, 15, 20]
    # Current read and historical cut use separate SQL sources but must agree today.
    for extra in ({}, {"as_of": today.isoformat()}):
        recon = await client.get(f"{root}/reconciliation", headers=headers, params=extra)
        assert recon.status_code == 200, recon.text
        data = recon.json()
        target = next(row for row in data["items"] if row["inventory_item_id"] == str(item_id))
        assert Decimal(target["physical_quantity"]) == Decimal(target["valued_quantity"]) == 30
        assert Decimal(target["inventory_value"]) == 500
        assert Decimal(target["quantity_difference"]) == 0
        assert data["reconciled_items"] >= 1
    before = await client.get(
        f"{root}/reconciliation",
        headers=headers,
        params={"as_of": (today - timedelta(days=1)).isoformat()},
    )
    assert before.status_code == 200, before.text
    assert before.json()["balance_known"] is False
    assert all(row["physical_quantity"] is None for row in before.json()["items"])
    literal = await client.get(f"{root}/reconciliation", headers=headers, params={"q": "%"})
    assert literal.json()["total"] == 0
    clean = await client.get(
        f"{root}/reconciliation", headers=headers, params={"difference_only": True}
    )
    assert all(row["quantity_difference"] != "0.000000" for row in clean.json()["items"])
    future = await client.get(
        f"{root}/reconciliation",
        headers=headers,
        params={"as_of": (today + timedelta(days=1)).isoformat()},
    )
    assert future.status_code == 422
    operator = await _operator_headers(client, company, warehouse)
    async with async_session_factory() as session:
        source = await session.get(Warehouse, warehouse)
        other = Warehouse(
            company_id=company,
            branch_id=source.branch_id,
            warehouse_category_id=source.warehouse_category_id,
            code="KDX-PRIVATE",
            name="Otra bodega de la misma sucursal",
            is_central=False,
        )
        session.add(other)
        await session.flush()
        unassigned_id = other.id
        await session.commit()
    # A branch grant alone does not authorize another warehouse in that branch.
    for path in ("physical", "physical/export", "items"):
        denied = await client.get(
            f"{root}/{path}",
            headers=operator,
            params={**physical_params, "warehouse_id": str(unassigned_id)},
        )
        assert denied.status_code == 403, denied.text
    for kind, query in (
        ("physical", physical_params),
        ("valuation", params),
        ("reconciliation", {"as_of": today.isoformat()}),
    ):
        for extension in ("pdf", "xlsx"):
            response = await client.get(
                f"{root}/{kind}/export", headers=headers, params={**query, "format": extension}
            )
            assert response.status_code == 200, response.text[:500]
            assert response.headers["cache-control"] == "no-store"
            assert f".{extension}" in response.headers["content-disposition"]
            (tmp_path / f"{kind}.{extension}").write_bytes(response.content)
            if extension == "pdf":
                assert response.content.startswith(b"%PDF-")
            else:
                book = load_workbook(BytesIO(response.content))
                if kind == "valuation":
                    sheet = book["Historial del costo"]
                    assert sheet.max_row == 4
                    assert sheet["J2"].value == 10  # First average remains historical.
                    assert sheet["J4"].value == pytest.approx(16.666667)
                    assert sheet["F3"].data_type == "n"
                elif kind == "physical":
                    sheet = book["Movimientos físicos"]
                    assert sheet.max_row == 2
                    assert sheet["G2"].value == 10
                    assert not any("costo" in str(cell.value).lower() for cell in sheet[1])
            allowed = await client.get(
                f"{root}/{kind}/export", headers=operator, params={**query, "format": extension}
            )
            assert allowed.status_code == (200 if kind == "physical" else 403), allowed.text[:200]
    foreign = await client.get(
        f"{root}/physical/export",
        headers=operator,
        params={**physical_params, "warehouse_id": str(uuid.uuid4())},
    )
    assert foreign.status_code == 404
    unauthenticated = await client.get(f"{root}/physical/export", params=physical_params)
    assert unauthenticated.status_code == 401
    # Reading and exporting do not append or rewrite the financial ledger.
    assert (await client.get(f"{root}/valuation", headers=headers, params=params)).json() == values
