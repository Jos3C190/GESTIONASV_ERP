"""Final acceptance through actual commercial receipts, placement and depletion."""

from datetime import datetime, timedelta
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from zoneinfo import ZoneInfo

import pytest
from app.infrastructure.db.session import async_session_factory
from app.infrastructure.models.inventory import InventoryBalanceModel, InventoryItemModel
from app.infrastructure.models.inventory_valuation import InventoryCostAverageModel
from app.infrastructure.models.organization import Location
from openpyxl import load_workbook
from sqlalchemy import func, select

from tests.e2e.test_inventory_preparation import _operator_headers, _review_unit
from tests.e2e.test_purchase_orders import (
    _headers,
    _receiving_location_id,
    _reference_ids,
    _sent_order,
)

pytestmark = pytest.mark.e2e
LOCAL_TZ = ZoneInfo("America/El_Salvador")


async def _balances(ctx):
    async with async_session_factory() as session:
        physical = await session.scalar(
            select(func.coalesce(func.sum(InventoryBalanceModel.quantity_base), 0)).where(
                InventoryBalanceModel.company_id == ctx.company,
                InventoryBalanceModel.inventory_item_id == ctx.item,
            )
        )
        balance = await session.get(InventoryCostAverageModel, (ctx.company, ctx.item))
        return (
            Decimal(physical),
            Decimal(balance.quantity_base) if balance else Decimal(0),
            Decimal(balance.inventory_value) if balance else Decimal(0),
        )


async def _receive_and_place(client, ctx, index, quantity, value, *, ordered_quantity=None):
    """Manual FOB is authoritative; no helper bypasses the business writes."""
    before = await _balances(ctx)
    order = await _sent_order(
        client,
        headers=ctx.admin,
        branch_id=ctx.branch,
        warehouse_id=ctx.warehouse,
        product_id=ctx.product,
        supplier_id=ctx.supplier,
        currency="USD",
        quantity=str(ordered_quantity or quantity),
    )
    created = await client.post(
        "/api/v1/purchases",
        headers=ctx.admin,
        json={
            "purchase_order_id": order["id"],
            "supplier_invoice_number": f"FAC-ACEPTACION-{index}",
            "supplier_invoice_date": ctx.today,
            "lines": [
                {
                    "purchase_order_detail_id": order["details"][0]["id"],
                    "quantity_received": str(quantity),
                }
            ],
        },
    )
    assert created.status_code == 201, created.text
    purchase = created.json()
    assert purchase["inventory_workflow"] == "placement_before_inventory"
    confirmed = await client.post(f"/api/v1/purchases/{purchase['id']}/receive", headers=ctx.admin)
    assert confirmed.status_code == 200, confirmed.text
    if ordered_quantity is not None and ordered_quantity > quantity:
        remaining = await client.get(
            f"/api/v1/purchase-orders/{order['id']}/receivable", headers=ctx.admin
        )
        assert remaining.status_code == 200, remaining.text
        assert Decimal(remaining.json()["lines"][0]["quantity_pending"]) == (
            ordered_quantity - quantity
        )
    else:
        received_order = await client.get(
            f"/api/v1/purchase-orders/{order['id']}", headers=ctx.admin
        )
        assert received_order.status_code == 200, received_order.text
        assert received_order.json()["status"] == "received"
    verification = {
        "lines": [
            {
                "purchase_detail_id": purchase["details"][0]["id"],
                "quantity_accepted": str(quantity),
                "inventory_location_id": str(ctx.receiving),
            }
        ]
    }
    premature = await client.post(
        f"/api/v1/purchases/{purchase['id']}/verify",
        headers=ctx.admin,
        json=verification,
    )
    assert premature.status_code == 409, premature.text
    # The three documented acquisition components are captured explicitly.
    retaceo = await client.post(
        "/api/v1/retaceos",
        headers=ctx.admin,
        json={
            "purchase_id": purchase["id"],
            "fob_lines": [
                {
                    "purchase_detail_id": purchase["details"][0]["id"],
                    "amount": str(value - 10),
                }
            ],
            "total_freight": "5",
            "total_dai": "5",
            "total_expenses": "0",
            "import_vat": "0",
            "notes": "FOB manual de factura de prueba; no reutilizar el precio de la oferta.",
        },
    )
    assert retaceo.status_code == 201, retaceo.text
    retaceo_id = retaceo.json()["id"]
    for action in ("calculate", "verify", "close"):
        closed = await client.post(f"/api/v1/retaceos/{retaceo_id}/{action}", headers=ctx.admin)
        assert closed.status_code == 200, closed.text
    assert Decimal(closed.json()["total_cost"]) == value
    assert closed.json()["status"] == "closed"
    assert all(not cost["documents"] for cost in closed.json()["costs"])
    availability = await client.get(
        f"/api/v1/purchase-orders/{order['id']}/actual-expenses/availability", headers=ctx.admin
    )
    assert availability.status_code == 200, availability.text
    assert availability.json()["can_record"] is False
    assert await _balances(ctx) == before  # Closing costs is not stock entry.
    comparison = await client.get(
        f"/api/v1/retaceos/{retaceo_id}/average-comparison", headers=ctx.admin
    )
    assert comparison.status_code == 200, comparison.text
    line = comparison.json()["lines"][0]
    assert Decimal(line["previous_quantity"]) == before[0]
    assert Decimal(line["previous_value"]) == before[2]
    assert Decimal(line["incoming_quantity"]) == quantity
    assert Decimal(line["incoming_value"]) == value
    assert Decimal(line["resulting_quantity"]) == before[0] + quantity
    expected_average = ((before[2] + value) / (before[0] + quantity)).quantize(Decimal("0.000001"))
    assert Decimal(line["resulting_average"]) == expected_average
    verified = await client.post(
        f"/api/v1/purchases/{purchase['id']}/verify",
        headers=ctx.admin,
        json=verification,
    )
    assert verified.status_code == 200, verified.text
    assert await _balances(ctx) == before
    queue = await client.get(
        "/api/v1/inventory/preparation",
        headers=ctx.operator,
        params={"warehouse_id": str(ctx.warehouse)},
    )
    assert queue.status_code == 200, queue.text
    assert len(queue.json()["items"]) == 1
    unit = queue.json()["items"][0]
    preview = await client.get(
        "/api/v1/inventory/labels/preview",
        headers=ctx.operator,
        params={"handling_unit_id": unit["id"]},
    )
    assert preview.status_code == 200, preview.text
    assert preview.json()["received_on"]
    assert preview.json()["receipt_lot_code"]
    await _review_unit(client, ctx.operator, unit)
    assert await _balances(ctx) == before  # Labels and measuring don't post stock.
    body = {
        "idempotency_key": f"final-placement-{index}",
        "quantity_base": str(quantity),
        "unit_scan": unit["code"],
        "location_scan": ctx.location_code,
    }
    placed = await client.post(
        f"/api/v1/inventory/preparation/{unit['id']}/place",
        headers=ctx.operator,
        json=body,
    )
    assert placed.status_code == 200, placed.text
    replay = await client.post(
        f"/api/v1/inventory/preparation/{unit['id']}/place",
        headers=ctx.operator,
        json=body,
    )
    assert replay.status_code == 200, replay.text
    assert replay.json() == placed.json()
    assert await _balances(ctx) == (before[0] + quantity, before[1] + quantity, before[2] + value)
    return unit


async def _sell_all(client, ctx, units):
    customer = await client.post(
        "/api/v1/sales/customers",
        headers=ctx.admin,
        json={"code": "FINAL-CLIENTE", "name": "Cliente de aceptación final"},
    )
    assert customer.status_code == 201, customer.text
    quote = await client.post(
        "/api/v1/sales/quotations",
        headers={**ctx.admin, "Idempotency-Key": "final-sale-quote"},
        json={
            "branch_id": str(ctx.branch),
            "customer_id": customer.json()["id"],
            "valid_until": (datetime.now(LOCAL_TZ).date() + timedelta(days=2)).isoformat(),
            "currency": "USD",
            "lines": [
                {
                    "product_id": ctx.product,
                    "unit_id": ctx.unit_id,
                    "quantity": "30",
                    "unit_price": "30",
                }
            ],
        },
    )
    assert quote.status_code == 201, quote.text
    confirmed = await client.post(
        f"/api/v1/sales/quotations/{quote.json()['id']}/confirm",
        headers=ctx.admin,
        json={"expected_version": quote.json()["version"]},
    )
    assert confirmed.status_code == 200, confirmed.text
    converted = await client.post(
        f"/api/v1/sales/quotations/{quote.json()['id']}/convert",
        headers={**ctx.admin, "Idempotency-Key": "final-convert"},
        json={"expected_version": confirmed.json()["version"]},
    )
    assert converted.status_code == 200, converted.text
    sale = await client.post(
        f"/api/v1/sales/{converted.json()['id']}/confirm",
        headers=ctx.admin,
        json={"expected_version": converted.json()["version"]},
    )
    assert sale.status_code == 200, sale.text
    assert await _balances(ctx) == (30, 30, 500)
    delivery = await client.post(
        f"/api/v1/sales/{sale.json()['id']}/deliveries",
        headers={**ctx.admin, "Idempotency-Key": "final-delivery"},
        json={
            "assignments": [
                {"sale_line_id": sale.json()["lines"][0]["id"], "handling_unit_id": unit["id"]}
                for unit in units
            ]
        },
    )
    assert delivery.status_code == 201, delivery.text
    assert await _balances(ctx) == (30, 30, 500)
    response = await client.post(
        f"/api/v1/sales/deliveries/{delivery.json()['id']}/confirm",
        headers={**ctx.admin, "Idempotency-Key": "final-delivery-confirm"},
        json={"expected_version": delivery.json()["version"]},
    )
    assert response.status_code == 200, response.text
    assert await _balances(ctx) == (0, 0, 0)


async def test_full_receipts_weighted_average_exhaustion_and_reports(e2e_client):
    client = e2e_client
    company, branch, warehouse, product, supplier, _ = await _reference_ids()
    admin = await _headers(client, company_id=company, username="final-admin", is_superuser=True)
    overview = await client.get("/api/v1/inventory/valuation", headers=admin)
    opening = await client.post(
        "/api/v1/inventory/valuation/opening",
        headers=admin,
        json={
            "currency": "USD",
            "reference": "Corte vacío para aceptación final",
            "approved": True,
            "lines": [],
            "snapshot_fingerprint": overview.json()["snapshot_fingerprint"],
        },
    )
    assert opening.status_code == 200, opening.text
    async with async_session_factory() as session:
        item = await session.scalar(
            select(InventoryItemModel).where(
                InventoryItemModel.company_id == company, InventoryItemModel.product_id == product
            )
        )
        destination = await session.scalar(
            select(Location).where(
                Location.warehouse_id == warehouse,
                Location.storage_eligible.is_(True),
                Location.is_active.is_(True),
                Location.lifecycle_status == "active",
                Location.location_type != "receiving",
            )
        )
        assert item is not None
        assert destination is not None
        ctx = SimpleNamespace(
            company=company,
            branch=branch,
            warehouse=warehouse,
            product=product,
            supplier=supplier,
            admin=admin,
            item=item.id,
            unit_id=item.base_unit_id,
            location_code=destination.code,
            today=datetime.now(LOCAL_TZ).date().isoformat(),
            receiving=await _receiving_location_id(warehouse),
        )
    ctx.operator = await _operator_headers(client, company, warehouse)
    units = [
        await _receive_and_place(client, ctx, i, Decimal(qty), Decimal(value))
        for i, (qty, value) in enumerate(((5, 50), (10, 150), (15, 300)), 1)
    ]
    root = "/api/v1/inventory/kardex"
    params = {"inventory_item_id": str(ctx.item), "period_from": ctx.today, "period_to": ctx.today}
    valued = await client.get(f"{root}/valuation", headers=admin, params=params)
    assert valued.status_code == 200, valued.text
    original_entries = valued.json()["entries"]
    assert [Decimal(e["average_unit_cost"]) for e in original_entries] == [
        Decimal("10"),
        Decimal("13.333333"),
        Decimal("16.666667"),
    ]
    assert [Decimal(e["balance_value"]) for e in original_entries] == [50, 200, 500]
    assert len({u["inventory_lot_detail_id"] for u in units}) == 3
    physical = await client.get(
        f"{root}/physical", headers=ctx.operator, params={**params, "warehouse_id": str(warehouse)}
    )
    assert physical.status_code == 200, physical.text
    assert physical.json()["total"] == 3
    assert Decimal(physical.json()["incoming_quantity"]) == 30
    assert all(row["document_reference"].startswith("COM-") for row in physical.json()["items"])
    await _sell_all(client, ctx, units)
    # With no remaining stock/value the next receipt takes its own landed cost.
    await _receive_and_place(client, ctx, 4, Decimal(4), Decimal(100), ordered_quantity=Decimal(6))
    final = await client.get(f"{root}/valuation", headers=admin, params=params)
    assert final.status_code == 200, final.text
    assert final.json()["entries"][:3] == original_entries
    assert Decimal(final.json()["entries"][-1]["average_unit_cost"]) == 25
    assert await _balances(ctx) == (4, 4, 100)
    reconciliation = await client.get(f"{root}/reconciliation", headers=admin)
    assert reconciliation.status_code == 200, reconciliation.text
    row = next(r for r in reconciliation.json()["items"] if r["inventory_item_id"] == str(ctx.item))
    assert Decimal(row["physical_quantity"]) == Decimal(row["valued_quantity"]) == 4
    assert Decimal(row["quantity_difference"]) == 0
    assert Decimal(row["inventory_value"]) == 100
    for extension in ("pdf", "xlsx"):
        exported = await client.get(
            f"{root}/valuation/export", headers=admin, params={**params, "format": extension}
        )
        assert exported.status_code == 200, exported.text[:300]
        if extension == "pdf":
            assert exported.content.startswith(b"%PDF-")
        else:
            sheet = load_workbook(BytesIO(exported.content))["Historial del costo"]
            assert sheet.max_row == 8  # 4 receipts plus 3 issued logistical units.
            assert sheet["J2"].value == 10
            assert sheet.cell(sheet.max_row, 10).value == 25
    denied = await client.get(f"{root}/valuation", headers=ctx.operator, params=params)
    assert denied.status_code == 403
    assert (
        await client.get(f"{root}/valuation", headers=admin, params=params)
    ).json() == final.json()
