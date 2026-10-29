"""Saved values, source quantities, numeric Excel cells and printable multipage PDFs."""

import re
import uuid
from dataclasses import replace
from datetime import date
from decimal import Decimal
from io import BytesIO

import pytest
from app.application.procurement_reports import consolidation_report, retaceo_report
from app.domain.entities.purchase_consolidation import (
    ConsolidationLine,
    ConsolidationSource,
    PurchaseConsolidation,
    PurchaseConsolidationStatus,
)
from app.domain.entities.tabular_report import ProcurementReportContext
from app.infrastructure.procurement_report_renderer import render_pdf, render_xlsx
from openpyxl import load_workbook

from tests.unit.test_retaceo_use_cases import FakeRetaceoRepository, _create, _ids, _purchase


@pytest.mark.asyncio
async def test_retaceo_export_preserves_saved_costs_and_separates_import_vat():
    purchase = _purchase(_ids())
    item = await _create(FakeRetaceoRepository(), purchase)
    context = ProcurementReportContext(
        "Grupo Lorena",
        purchase_code=purchase.code,
        product_labels={11: ("PRD-11", "Azúcar Glass"), 12: ("PRD-12", "Café")},
    )
    report = retaceo_report(item, context)
    assert report.sections[0].rows[0][4] == item.details[0].cost_fob
    assert report.sections[0].rows[0][9] == item.details[0].unit_cost
    summary = dict(zip(report.sections[1].columns, report.sections[1].rows[0], strict=True))
    assert summary["Costo aterrizado"] == item.total_cost
    assert summary["IVA de importación (fuera del costo)"] == item.import_vat
    assert (
        item.total_cost
        == item.total_fob + item.total_freight + item.total_expenses + item.total_dai
    )
    book = load_workbook(BytesIO(render_xlsx(report)))
    sheet = book["Costos por producto"]
    assert sheet["E2"].data_type == "n"
    assert Decimal(str(sheet["E2"].value)) == item.details[0].cost_fob
    assert sheet["J2"].number_format == "#,##0.000000"
    assert sheet.freeze_panes == "A2"
    assert book["Resumen del retaceo"]["E2"].value == item.total_cost
    pdf = render_pdf(report)
    assert pdf.startswith(b"%PDF-")
    assert len(pdf) > 1500
    # Enough rows to exercise repeated headers and PDF pagination.
    large = replace(
        report, sections=(replace(report.sections[0], rows=report.sections[0].rows * 100),)
    )
    assert len(re.findall(rb"/Type /Page\b", render_pdf(large))) > 1


@pytest.mark.asyncio
async def test_excel_keeps_untrusted_labels_as_text_and_pdf_escapes_markup():
    purchase = _purchase(_ids())
    item = await _create(FakeRetaceoRepository(), purchase)
    malicious = '=HYPERLINK("https://example.invalid", "Abrir")'
    report = retaceo_report(
        item,
        ProcurementReportContext(
            "Grupo Lorena", product_labels={11: (malicious, "<b>Azúcar & café</b>\x00")}
        ),
    )
    book = load_workbook(BytesIO(render_xlsx(report)), data_only=False)
    assert book["Costos por producto"]["A2"].value == malicious
    assert book["Costos por producto"]["A2"].data_type == "s"
    assert render_pdf(report).startswith(b"%PDF-")


def test_general_consolidation_aggregates_branches_per_product_and_keeps_source_lineage():
    source_ids = [uuid.uuid4() for _ in range(7)]
    source = ConsolidationSource(
        *source_ids[:3],
        "SCR-001",
        source_ids[3],
        "Roosevelt",
        source_ids[4],
        "Bodega",
        "Rodrigo",
        None,
        Decimal("20"),
        Decimal("18"),
        "Ajuste justificado",
    )
    second = replace(
        source,
        id=uuid.uuid4(),
        request_code="SCR-002",
        branch_name="Gotera",
        requested_quantity=Decimal("30"),
        included_quantity=Decimal("30"),
        reason=None,
    )
    line = ConsolidationLine(
        uuid.uuid4(),
        4,
        "PRD-AZU-004",
        "Azúcar",
        2,
        "Caja",
        Decimal("2"),
        "Reserva central",
        (source, second),
    )
    item = PurchaseConsolidation(
        uuid.uuid4(),
        uuid.uuid4(),
        "CON-001",
        date(2026, 10, 1),
        date(2026, 10, 27),
        PurchaseConsolidationStatus.FROZEN,
        uuid.uuid4(),
        None,
        None,
        uuid.uuid4(),
        "Matriz",
        None,
        (line,),
    )
    report = consolidation_report(item, ProcurementReportContext("Grupo Lorena"))
    assert report.sections[0].rows[0][3:] == (Decimal("48"), Decimal("2"), Decimal("50"))
    assert report.sections[1].rows[0][2] == "SCR-001"
    assert report.sections[1].rows[0][4:6] == (Decimal("20"), Decimal("18"))
    assert render_pdf(report).startswith(b"%PDF-")
