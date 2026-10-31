"""Reports retain Decimal values, readable references, and complete query results."""

from datetime import UTC, date, datetime
from decimal import Decimal
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from app.api.v1.routers import inventory_kardex as read_routes
from app.api.v1.routers import inventory_kardex_reports as routes
from app.application.inventory.kardex_reports import kardex_report
from app.infrastructure.procurement_report_renderer import render_pdf, render_xlsx
from openpyxl import load_workbook


def _statement(kind):
    context = {
        "company_name": "Grupo Lorena",
        "nit": "NIT-DEMO",
        "nrc": "NRC-DEMO",
        "product_name": "Azúcar Glass",
        "sku": "PRD-AZU-004",
        "base_unit_symbol": "cj",
    }
    data = {
        "context": context,
        "period_from": date(2026, 10, 1),
        "period_to": date(2026, 10, 29),
        "opening_quantity": Decimal("0"),
        "opening_value": Decimal("0"),
        "closing_quantity": Decimal("5"),
        "closing_value": Decimal("50"),
        "period_value_change": Decimal("50"),
        "balance_basis": "Corte aprobado",
        "incoming_quantity": Decimal("5"),
        "outgoing_quantity": Decimal("0"),
        "setting": {"currency": "USD"},
        "as_of": date(2026, 10, 29),
        "generated_at": datetime(2026, 10, 29, 23, tzinfo=UTC),
        "currency": "USD",
        "inventory_value": Decimal("50"),
        "reconciled_items": 1,
        "discrepant_items": 0,
        "unvalued_items": 0,
    }
    if kind == "physical":
        data["items"] = [
            {
                "posted_at": data["generated_at"],
                "movement_type": "receipt",
                "document_reference": "COM-00027",
                "receipt_lot_code": "LOT-COM-00027",
                "to_location_code": "A01-R01-N01-P01",
                "operator": "rodrigo",
                "incoming_quantity": Decimal("5"),
                "outgoing_quantity": Decimal("0"),
                "balance_after": Decimal("5"),
            }
        ]
    elif kind == "valuation":
        data["entries"] = [
            {
                "accounting_date": date(2026, 10, 29),
                "kind": "receipt",
                "source_reference": "technical-id",
                "document_reference": "COM-00027",
                "receipt_lot_code": "LOT-COM-00027",
                "operator": "rodrigo",
                "quantity_base": Decimal("5"),
                "applied_unit_cost": Decimal("10"),
                "inventory_delta": Decimal("50"),
                "balance_quantity": Decimal("5"),
                "balance_value": Decimal("50"),
                "average_unit_cost": Decimal("10"),
            }
        ]
    else:
        data["items"] = [
            {
                "sku": "PRD-AZU-004",
                "product_name": "Azúcar Glass",
                "base_unit_symbol": "cj",
                "physical_quantity": Decimal("5"),
                "valued_quantity": Decimal("5"),
                "quantity_difference": Decimal("0"),
                "average_unit_cost": Decimal("10"),
                "inventory_value": Decimal("50"),
            }
        ]
    return data


@pytest.mark.parametrize("kind", ["physical", "valuation", "reconciliation"])
def test_kardex_report_pdf_and_numeric_excel(kind):
    report = kardex_report(kind, _statement(kind), {})
    assert ("NIT", "NIT-DEMO") in report.metadata
    assert ("NRC", "NRC-DEMO") in report.metadata
    assert render_pdf(report).startswith(b"%PDF-")
    book = load_workbook(BytesIO(render_xlsx(report)))
    if kind == "valuation":
        sheet = book["Historial del costo"]
        assert sheet["C2"].value == "COM-00027\nLote: LOT-COM-00027"
        assert sheet["F2"].data_type == "n"
        assert sheet["J2"].value == 10
        assert sheet["F2"].number_format == "#,##0.000000"
    elif kind == "physical":
        assert book["Movimientos físicos"]["G2"].value == 5
    else:
        assert book["Conciliación por producto"]["F2"].value == 0


def test_unknown_balances_and_value_only_adjustments_do_not_invent_numbers():
    data = _statement("valuation")
    data["opening_quantity"] = None
    data["entries"][0].update(
        kind="cost_adjustment", quantity_base=Decimal("0"), applied_unit_cost=None
    )
    report = kardex_report("valuation", data, {})
    assert report.sections[0].rows[0][0] == "No reconstruible"
    assert report.sections[1].rows[0][4:6] == (Decimal("0"), "No aplica")


@pytest.mark.asyncio
async def test_export_fetches_every_page_and_rejects_oversized_instead_of_truncating(monkeypatch):
    first = _statement("physical")
    first["items"] *= 200
    fetch = AsyncMock(
        side_effect=[
            SimpleNamespace(total=201, model_dump=lambda: first),
            SimpleNamespace(model_dump=lambda: {"items": [first["items"][0]]}),
        ]
    )
    response = await routes._export(fetch, "physical", "xlsx", {})
    assert fetch.await_args_list[1].kwargs == {"page": 2, "page_size": 200}
    book = load_workbook(BytesIO(response.body))
    assert book["Movimientos físicos"].max_row == 202
    fetch = AsyncMock(return_value=SimpleNamespace(total=5001))
    with pytest.raises(routes.HTTPException) as error:
        await routes._export(fetch, "physical", "pdf", {})
    assert error.value.status_code == 422
    assert fetch.await_count == 1


@pytest.mark.asyncio
async def test_financial_scope_rejects_a_warehouse_restricted_membership(monkeypatch):
    monkeypatch.setattr(read_routes, "require_company_wide_scope", AsyncMock())
    monkeypatch.setattr(read_routes, "assigned_warehouse_ids", AsyncMock(return_value=[]))
    with pytest.raises(routes.HTTPException) as error:
        await read_routes._authorize_company_values(None, None, None)
    assert error.value.status_code == 403
