"""Read-only independent arithmetic checks of the supplied workbook."""
import hashlib
import json
import re
from decimal import Decimal, ROUND_HALF_UP, getcontext
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

getcontext().prec = 36
root = Path(__file__).parent
source = Path(r"C:\Users\rodri\Downloads\Retaceo (4).xlsx")
original_hash = hashlib.sha256(source.read_bytes()).hexdigest()
formulas = load_workbook(source, data_only=False)
cached = load_workbook(source, data_only=True)
D = lambda value: Decimal(str(value or 0))
money = lambda value: str(value.quantize(Decimal(".01"), rounding=ROUND_HALF_UP))


def value(sheet, address):
    raw = sheet[address].value
    if not isinstance(raw, str) or not raw.startswith("="):
        return D(raw) if raw is None or isinstance(raw, (int, float)) else raw
    expr = raw[1:].replace("$", "")
    match = re.fullmatch(r"SUM\(([A-Z]+\d+):([A-Z]+\d+)\)", expr)
    if match:
        left, top, right, bottom = range_boundaries(":".join(match.groups()))
        return sum((value(sheet, sheet.cell(row, col).coordinate)
                    for row in range(top, bottom + 1)
                    for col in range(left, right + 1)), Decimal(0))
    match = re.fullmatch(r"([A-Z]+\d+)\*IF\(([A-Z]+\d+)>0,([A-Z]+\d+)/([A-Z]+\d+),0\)", expr)
    if match:
        line, guard, numerator, denominator = match.groups()
        return value(sheet, line) * (value(sheet, numerator) / value(sheet, denominator)
                                     if value(sheet, guard) > 0 else Decimal(0))
    match = re.fullmatch(r"([A-Z]+\d+)\*([A-Z]+\d+)", expr)
    if match:
        return value(sheet, match[1]) * value(sheet, match[2])
    match = re.fullmatch(r'"([^"]*)"&TEXT\(IF\(([A-Z]+\d+)>0,([A-Z]+\d+)/([A-Z]+\d+),0\),"0.00%"\)', expr)
    if match:
        label, guard, numerator, denominator = match.groups()
        ratio = value(sheet, numerator) / value(sheet, denominator) if value(sheet, guard) > 0 else Decimal(0)
        return label + money(ratio * 100) + "%"
    raise ValueError(f"Unsupported formula: {sheet.title}!{address}: {raw}")


checks = []
for sheet in formulas:
    for row in sheet:
        for cell in row:
            if cell.data_type != "f":
                continue
            computed = value(sheet, cell.coordinate)
            saved = cached[sheet.title][cell.coordinate].value
            matches = abs(computed - D(saved)) < Decimal(".00000001") if isinstance(computed, Decimal) else computed == saved
            checks.append(dict(sheet=sheet.title, cell=cell.coordinate, formula=cell.value,
                               independently_computed=str(computed), saved=saved, matches=matches))

sheet = formulas["Retaceo"]
data = formulas["Datos"]
fob = value(sheet, "G10")
freight = value(sheet, "G11")
expenses = value(sheet, "I10")
dai = value(sheet, "I11")
insurance = value(data, "D5")
included_other = sum(value(data, f"D{row}") for row in range(9, 19))
corrected_total = fob + freight + expenses + dai + insurance
insurance_scenario = []
for row in (15, 16):
    line_fob = value(sheet, f"E{row}")
    added_insurance = line_fob / fob * insurance
    line_total = line_fob + line_fob / fob * (freight + expenses + dai + insurance)
    insurance_scenario.append(dict(product=sheet[f"C{row}"].value,
                                   insurance_allocation=money(added_insurance),
                                   corrected_line_cost=money(line_total),
                                   original_line_cost=money(value(sheet, f"I{row}"))))

comparison = []
for provider, prices, shown in (("A", (800, 15, 20), 4450), ("B", (780, 13, 19), 4240), ("C", (825, 14, 18), 4405)):
    computed = sum(q * p for q, p in zip((5, 10, 10), prices))
    comparison.append(dict(provider=provider, computed=computed, document=shown, difference=shown-computed))

result = dict(
    source=str(source), sha256=original_hash,
    formula_count=len(checks), all_cached_results_match=all(c["matches"] for c in checks),
    formula_checks=checks,
    reconciliation=dict(fob=money(fob), freight=money(freight), insurance=money(insurance),
                        listed_other_costs=money(included_other), dai=money(dai),
                        retaceo_total=money(value(sheet, "I19")), inventory_total=money(value(data, "L22")),
                        difference=money(value(data, "L22")-value(sheet, "I19")),
                        total_with_insurance=money(corrected_total), cif=money(value(data, "I2")),
                        workbook_import_vat=money(value(data, "D8"))),
    insurance_scenario=insurance_scenario,
    quantity_two_scenario=dict(line_fob_unchanged=money(value(sheet, "E15")), quantity=2,
                              workbook_shown_unit_cost=money(value(sheet, "I15")),
                              actual_unit_cost=money(value(sheet, "I15")/2)),
    extra_fob_5000_without_header_update=dict(
        line_fob_total=money(fob+5000), denominator_still=money(fob),
        actual_costs_to_allocate=money(freight+expenses+dai),
        workbook_allocated=money((fob+5000)/fob*(freight+expenses+dai))),
    pdf_comparison_example=comparison,
    source_unchanged=original_hash == hashlib.sha256(source.read_bytes()).hexdigest(),
)
(root / "math-audit.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
print(json.dumps({k: v for k, v in result.items() if k != "formula_checks"}, ensure_ascii=False, indent=2))
assert result["all_cached_results_match"]
assert result["source_unchanged"]
