from pathlib import Path
import json
from pypdf import PdfReader
from openpyxl import load_workbook

BASE = Path(__file__).resolve().parent
PDF = Path(r'C:\Users\rodri\Downloads\MiniERP ERS v0.9 (2).pdf')
XLSX = Path(r'C:\Users\rodri\Downloads\Retaceo (4).xlsx')

reader = PdfReader(PDF)
pages = []
for i, page in enumerate(reader.pages, 1):
    text = page.extract_text(extraction_mode='layout')
    pages.append({'page': i, 'text': text, 'images': len(page.images)})
    (BASE / f'page-{i:02d}.txt').write_text(text, encoding='utf-8')
(BASE / 'pdf-text.txt').write_text('\n\n'.join(f'===== PAGINA {p["page"]} =====\n{p["text"]}' for p in pages), encoding='utf-8')
print('PDF', json.dumps({'pages': len(pages), 'metadata': dict(reader.metadata or {}), 'page_sizes': [(float(p.mediabox.width), float(p.mediabox.height)) for p in reader.pages]}, ensure_ascii=False, default=str))
print('PDF page text lengths', [(p['page'], len(p['text']), p['images']) for p in pages])

wb = load_workbook(XLSX, data_only=False)
cached = load_workbook(XLSX, data_only=True)
report = {'properties': str(wb.properties), 'calculation': str(wb.calculation), 'defined_names': {k:str(v) for k,v in wb.defined_names.items()}, 'sheets': []}
lines = []
for ws in wb:
    cells = []
    for row in ws:
        row_cells = []
        for cell in row:
            if cell.value is None:
                continue
            entry = {'cell':cell.coordinate, 'type':cell.data_type, 'value':cell.value, 'cached':cached[ws.title][cell.coordinate].value, 'format':cell.number_format, 'comment':cell.comment.text if cell.comment else None, 'hyperlink':cell.hyperlink.target if cell.hyperlink else None}
            cells.append(entry)
            row_cells.append(f'{cell.coordinate}: {cell.value!s}' + (f' => {entry["cached"]}' if cell.data_type == 'f' else ''))
        if row_cells: lines.append(' | '.join(row_cells))
    meta = {'sheet':ws.title, 'dimensions':ws.calculate_dimension(), 'state':ws.sheet_state, 'merges':list(map(str,ws.merged_cells.ranges)), 'tables':list(ws.tables), 'validations':[str(x) for x in ws.data_validations.dataValidation], 'protection':str(ws.protection), 'hidden_rows':[k for k,v in ws.row_dimensions.items() if v.hidden], 'hidden_columns':[k for k,v in ws.column_dimensions.items() if v.hidden], 'print_area':str(ws.print_area), 'cells':cells}
    report['sheets'].append(meta)
    lines.insert(len(lines)-sum(1 for row in ws if any(c.value is not None for c in row)), f'===== HOJA {ws.title}: {ws.calculate_dimension()} =====')
    print('SHEET', json.dumps({k:v for k,v in meta.items() if k not in ['cells','protection']},ensure_ascii=False,default=str))
    print('COUNTS',len(cells),'cells;',sum(c['type']=='f' for c in cells),'formulas')
(BASE/'workbook.json').write_text(json.dumps(report,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
(BASE/'workbook.txt').write_text('\n'.join(lines),encoding='utf-8')
